# -*- coding: utf-8 -*-
"""按 layout.json 生成单页 PPTX（图形=图片，文本=原生文本框），并输出 PIL 预览与原图差异。

用法：
  python layout_to_pptx.py --layout layout.json --out deck.pptx [--preview preview.png] [--ref 原图.png]

layout.json 结构见 SKILL.md。所有相对路径按 layout.json 所在目录解析。
"""
import argparse
import json
import os
import sys

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.stdout.reconfigure(encoding="utf-8")
EMU_PER_PT = 12700


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--layout", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--preview", default=None, help="输出 PIL 预览 PNG（整页）")
    ap.add_argument("--ref", default=None, help="原图路径；给了就输出并排图与差异统计")
    ap.add_argument("--ref-top", type=int, default=None, help="原图在整页中的上边距（默认自动算）")
    ap.add_argument("--ea-font", default="微软雅黑", help="中文字体名（写入 a:ea 与 a:cs）")
    ap.add_argument("--lat-font", default="微软雅黑", help="拉丁字体名（写入 a:latin）")
    ap.add_argument("--font-file", default=r"C:\Windows\Fonts\msyh.ttc")
    ap.add_argument("--font-file-bold", default=r"C:\Windows\Fonts\msyhbd.ttc")
    args = ap.parse_args()

    base = os.path.dirname(os.path.abspath(args.layout))
    L = json.load(open(args.layout, encoding="utf-8"))

    def R(p):
        return p if os.path.isabs(p) else os.path.join(base, p)

    CW, CH = L["canvas"]
    w_pt, h_pt = L["slide"]["w_pt"], L["slide"]["h_pt"]
    SLIDE_W, SLIDE_H = int(w_pt * EMU_PER_PT), int(h_pt * EMU_PER_PT)
    S = SLIDE_W / CW                       # EMU per px
    PX2PT = w_pt / CW
    TOP_EMU = (SLIDE_H - CH * S) / 2.0     # 内容区上边距（EMU）
    TOP_PX = int(round(TOP_EMU / S))       # 同上，换算成原图像素
    print(f"canvas {CW}x{CH}  slide {w_pt}x{h_pt}pt  1px={PX2PT:.4f}pt  内容区上边距={TOP_PX}px")

    def measure(text, size_pt, bold):
        f = ImageFont.truetype(args.font_file_bold if bold else args.font_file,
                               int(round(size_pt / PX2PT)), index=0)
        return f.getbbox(text, anchor="la"), f

    # ---- 解算文本位置 ----
    for t in L["texts"]:
        bb, _ = measure(t["text"], t["size_pt"], t.get("bold", False))
        bx1, by1, bx2, by2 = bb
        if t.get("mode") == "center":
            ax1, ay1, ax2, ay2 = t["anchor_box"]
            t["_left"] = ax1
            t["_w"] = ax2 - ax1
            t["_top"] = (ay1 + ay2) / 2.0 - (by1 + by2) / 2.0
            t["_center"] = True
        else:
            x1, y1, x2, y2 = t["ink_box"]
            t["_left"] = x1 - bx1
            t["_top"] = (y1 + y2) / 2.0 - (by1 + by2) / 2.0
            t["_w"] = (x2 - x1) + (bx2 - bx1) + 12
            t["_center"] = False
        t["_size_px"] = int(round(t["size_pt"] / PX2PT))

    # ---- 写 PPTX ----
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sh = slide.shapes

    def E(v):
        return Emu(int(round(v)))

    if L.get("background"):
        p = sh.add_picture(R(L["background"]), 0, 0, E(SLIDE_W), E(SLIDE_H))
        p.name = "背景"
    for im in L["images"]:
        x1, y1, x2, y2 = im["box"]
        p = sh.add_picture(R(im["path"]), E(x1 * S), E(TOP_EMU + y1 * S),
                           E((x2 - x1) * S), E((y2 - y1) * S))
        p.name = os.path.basename(im["path"]).rsplit(".", 1)[0]

    for i, t in enumerate(L["texts"], 1):
        th = max(18.0, t["size_pt"] / PX2PT * 1.9)
        tb = sh.add_textbox(E(t["_left"] * S), E(TOP_EMU + t["_top"] * S), E(t["_w"] * S), E(th * S))
        tb.name = f"{i:02d}-{t['text'][:10]}"
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER if t["_center"] else PP_ALIGN.LEFT
        r = para.add_run()
        r.text = t["text"]
        f = r.font
        f.size = Pt(t["size_pt"])
        f.bold = bool(t.get("bold"))
        f.name = args.lat_font
        f.color.rgb = RGBColor(*t["color"])
        rPr = f._element                      # Font 直接代理 rPr
        for tag in ("a:ea", "a:cs"):
            rPr.append(rPr.makeelement(qn(tag), {"typeface": args.ea_font}))
        if t.get("spc"):
            rPr.set("spc", str(int(t["spc"])))

    notes = L.get("notes")
    if notes:
        slide.notes_slide.notes_text_frame.text = notes

    prs.save(args.out)
    chk = Presentation(args.out)
    sl = chk.slides[0]
    print(f"已写出 {args.out}")
    print(f"  形状 {len(sl.shapes)}：图片 {sum(1 for s in sl.shapes if s.shape_type == 13)} "
          f"文本框 {sum(1 for s in sl.shapes if s.has_text_frame)}；"
          f"文本 {sum(len(s.text_frame.text) for s in sl.shapes if s.has_text_frame)} 字")

    # ---- PIL 预览 ----
    if args.preview:
        fh = int(round(SLIDE_H / S))
        if L.get("background"):
            canvas = np.array(Image.open(R(L["background"])).convert("RGB"))
        else:
            canvas = np.full((fh, CW, 3), 255, dtype=np.uint8)
        pv = canvas.copy()
        for im in L["images"]:
            x1, y1, x2, y2 = im["box"]
            pv[TOP_PX + y1:TOP_PX + y2, x1:x2] = np.array(Image.open(R(im["path"])).convert("RGB"))
        pimg = Image.fromarray(pv)
        d = ImageDraw.Draw(pimg)
        for t in L["texts"]:
            f = ImageFont.truetype(args.font_file_bold if t.get("bold") else args.font_file,
                                   t["_size_px"], index=0)
            spc_px = (t["spc"] / 100.0 / PX2PT) if t.get("spc") else 0.0
            col = tuple(t["color"])
            y = TOP_PX + t["_top"]
            if t["_center"]:
                cx = (t["anchor_box"][0] + t["anchor_box"][2]) / 2.0
                if spc_px:
                    x = cx - (sum(f.getlength(c) for c in t["text"]) + spc_px * (len(t["text"]) - 1)) / 2
                    for c in t["text"]:
                        d.text((x, y), c, font=f, fill=col, anchor="la")
                        x += f.getlength(c) + spc_px
                else:
                    d.text((cx, y), t["text"], font=f, fill=col, anchor="ma")
            else:
                if spc_px:
                    x = t["_left"]
                    for c in t["text"]:
                        d.text((x, y), c, font=f, fill=col, anchor="la")
                        x += f.getlength(c) + spc_px
                else:
                    d.text((t["_left"], y), t["text"], font=f, fill=col, anchor="la")
        pimg.save(args.preview)
        print(f"预览 {args.preview}")

        if args.ref:
            ref = np.array(Image.open(args.ref).convert("RGB"))
            top = args.ref_top if args.ref_top is not None else TOP_PX
            a = ref.astype(np.int16)
            b = np.array(pimg).astype(np.int16)[top:top + CH]
            diff = np.abs(a - b).max(axis=2)
            print(f"\n与原图差异：mean={diff.mean():.2f} 中位={np.median(diff):.0f} "
                  f">30={100.0*(diff>30).mean():.2f}%  >60={100.0*(diff>60).mean():.2f}%")
            GS = 40
            rows = sorted(((float(diff[y:y+GS, x:x+GS].mean()), x, y)
                           for y in range(0, CH, GS) for x in range(0, CW, GS)), reverse=True)
            print("差异最大 8 块（文字区差异=字形差异；图形/背景区差异=真问题）：")
            for mn, x, y in rows[:8]:
                print(f"   x={x:4d} y={y:3d} mean={mn:6.1f}")
            cd = os.path.splitext(args.preview)[0] + "-compare.png"
            cmpi = Image.new("RGB", (CW, CH * 2 + 20), (255, 255, 255))
            cmpi.paste(Image.fromarray(ref), (0, 0))
            cmpi.paste(pimg.crop((0, top, CW, top + CH)), (0, CH + 20))
            cmpi.save(cd)
            print(f"上下对比图 {cd}")


if __name__ == "__main__":
    main()
