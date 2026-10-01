# -*- coding: utf-8 -*-
"""把 **PowerPoint 原生表格** 插入到已生成的 PPTX 里（复刻链路的「表格原生化」步骤）。

为什么单独一个脚本：
  `layout_to_pptx.py` 的产物是「位图底 + 位图切片 + 原生文本框」。表格区域如果也走位图，
  文字就改不动。本脚本把该区域替换成真正的 `<a:tbl>` —— 单元格可选中、可改字、可调行列尺寸、
  可合并单元格，且带显式边框/填充/对齐（不依赖主题表格样式）。

z-order：默认把表格插到「最后一个图形之后、第一个文本框之前」(`--z under-text`)，
  这样表格是原生的、文字层仍在其上，与 layout_to_pptx 的绘制次序一致。`--z top` 则置顶。

用法（build）：
  python add_tables.py --layout layout.json --tables tables.json \
      --pptx deck.pptx --out deck-table.pptx --verify

用法（probe，只读，从原图探测网格线并给出 col_widths / row_heights 建议）：
  python add_tables.py probe --ref 原图.png --box 120,300,1100,760

tables.json 结构见 ../SKILL.md「原生表格」一节。
"""
import argparse
import json
import os
import sys
from collections import Counter

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.stdout.reconfigure(encoding="utf-8")

EMU_PER_PT = 12700.0

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_VANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_SIDE_ORDER = ("L", "R", "T", "B")


def hx(rgb):
    return "%02X%02X%02X" % (int(rgb[0]), int(rgb[1]), int(rgb[2]))


def col(v, default):
    """'#RRGGBB' / [r,g,b] / None -> [r,g,b]"""
    if v is None:
        return list(default)
    if isinstance(v, str):
        s = v.strip().lstrip("#")
        return [int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)]
    return [int(x) for x in v]


# ---------------------------------------------------------------- 单元格边框


def _mk_line(tcPr, side, rgb, w_pt, style="solid"):
    tag = qn("a:ln" + side)
    ln = tcPr.makeelement(tag, {
        "w": str(int(round(w_pt * EMU_PER_PT))),
        "cap": "flat", "cmpd": "sng", "algn": "ctr",
    })
    if rgb is None:
        ln.append(ln.makeelement(qn("a:noFill"), {}))
    else:
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": hx(rgb)}))
        ln.append(fill)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": style}))
    return ln


def apply_borders(tc, rgb, w_pt, sides, style="solid"):
    """画/清单元格四边。sides 是 'LRTB' 的子集；rgb=None 表示整块去线。

    a:lnL,lnR,lnT,lnB 必须是 tcPr 的**前四个**子元素（CT_TableCellProperties 序列），
    所以倒序 insert(0, ...)。
    """
    tcPr = tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("a:lnL")) + tcPr.findall(qn("a:lnR")) \
            + tcPr.findall(qn("a:lnT")) + tcPr.findall(qn("a:lnB")):
        tcPr.remove(e)
    made = []
    for s in _SIDE_ORDER:
        if rgb is not None and s in sides:
            made.append(_mk_line(tcPr, s, rgb, w_pt, style))
        else:
            made.append(_mk_line(tcPr, s, None, 0))
    for el in reversed(made):
        tcPr.insert(0, el)


# ---------------------------------------------------------------- 单元格式样


def style_cell(cell, spec, st, write_text=True):
    """把一条 cell 定义（已合并全局 style）落到 python-pptx 的 cell 上。"""
    spec = spec or {}

    fill = spec.get("fill")
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*col(fill, [255, 255, 255]))

    mt, mr, mb, ml = st.get("margin_pt", [2.0, 3.0, 2.0, 3.0])
    cell.margin_top = Emu(int(mt * EMU_PER_PT))
    cell.margin_bottom = Emu(int(mb * EMU_PER_PT))
    cell.margin_left = Emu(int(ml * EMU_PER_PT))
    cell.margin_right = Emu(int(mr * EMU_PER_PT))
    cell.vertical_anchor = _VANCHOR[spec.get("valign") or st.get("valign", "middle")]

    if spec.get("no_border"):
        apply_borders(cell._tc, None, 0, "")
    else:
        b = spec.get("border") or st.get("border") or {}
        apply_borders(cell._tc,
                      col(b.get("color", st.get("border_color", [200, 210, 225])), [200, 210, 225]),
                      float(b.get("w_pt", st.get("border_w_pt", 0.75))),
                      b.get("sides", st.get("border_sides", "LRTB")),
                      b.get("style", "solid"))

    if not write_text:
        return

    text = spec.get("text", "")
    if isinstance(text, (int, float)):
        text = str(text)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text).split("\n") if text != "" else [""]
    paras = [tf.paragraphs[0]]
    for _ in range(len(lines) - 1):
        paras.append(tf.add_paragraph())

    fs = spec.get("size_pt", st.get("size_pt", 12))
    bold = spec.get("bold", st.get("bold", False))
    color = spec.get("color", st.get("color", [25, 25, 25]))
    align = _ALIGN[spec.get("align") or st.get("align", "center")]
    spc = spec.get("spc")
    ea = st.get("font", "微软雅黑")
    latin = st.get("latin_font", ea)

    for para, line in zip(paras, lines):
        para.alignment = align
        para.line_spacing = 1.0
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        r = para.add_run()
        r.text = line
        f = r.font
        f.size = Pt(float(fs))
        f.bold = bool(bold)
        f.name = latin
        f.color.rgb = RGBColor(*col(color, [25, 25, 25]))
        rPr = f._element                       # Font 直接代理 rPr
        for tag in ("a:ea", "a:cs"):
            rPr.append(rPr.makeelement(qn(tag), {"typeface": ea}))
        if spc:
            rPr.set("spc", str(int(spc)))


# ---------------------------------------------------------------- 表格构建


def _resolve_dims(vals, total, n, what, name):
    if not vals:
        return [total / n] * n
    if len(vals) != n:
        raise ValueError(f"表格 {name} 的 {what} 有 {len(vals)} 项，应为 {n} 项")
    s = float(sum(vals))
    if s <= 0:
        raise ValueError(f"表格 {name} 的 {what} 之和必须为正")
    if s <= 1.5:                                   # 判定为比例
        return [total * v / s for v in vals]
    if abs(s - total) > 1.0:
        print(f"  ! {what} 之和 {s:.1f}px 与 box 的 {total}px 不符，按比例缩放")
    return [total * v / s for v in vals]


def build_table(shapes, spec, S, TOP_EMU, st):
    """建一个原生表格，返回它的 <p:graphicFrame> 元素（已挂在 spTree 上）。

    单元格用**占用表**（occupancy grid）解算：`cells[r]` 只列该行的「起始单元格」，
    被 `colspan` / `rowspan` 吃掉的位置**不需要写占位**，由算法自动跳过并校验无重叠、无空洞。
    """
    name = spec.get("name") or "table"
    x1, y1, x2, y2 = spec["box"]
    wpx, hpx = x2 - x1, y2 - y1

    cells = [[(c or {}) for c in row] for row in spec["cells"]]
    nrow = len(cells)
    ncol = max(sum(int((c or {}).get("colspan", 1) or 1) for c in row) for row in cells)

    occ = [[None] * ncol for _ in range(nrow)]     # (r,c) -> 起始格 (r0,c0)
    origins = []                                   # (r0, c0, cell, r1, c1)
    for r in range(nrow):
        c = 0
        for cell in cells[r]:
            while c < ncol and occ[r][c] is not None:
                c += 1
            if c >= ncol:
                raise ValueError(f"表格 {name} 第 {r+1} 行的单元格超过 {ncol} 列，请检查 colspan")
            cs = int(cell.get("colspan", 1) or 1)
            rs = int(cell.get("rowspan", 1) or 1)
            if c + cs > ncol:
                raise ValueError(f"表格 {name} 第 {r+1} 行的 colspan={cs} 越出 {ncol} 列")
            r1, c1 = min(nrow - 1, r + rs - 1), c + cs - 1
            if rs > 1 and r + rs - 1 > nrow - 1:
                print(f"  ! 表格 {name} ({r},{c}) 的 rowspan={rs} 越出末行，"
                      f"已截断为 {r1 - r + 1}（是否少写了一行？）")
            for rr in range(r, r1 + 1):
                for cc in range(c, c1 + 1):
                    if occ[rr][cc] is not None:
                        raise ValueError(
                            f"表格 {name} 的合并区在 ({rr},{cc}) 发生重叠"
                            f"（与起始格 {occ[rr][cc]} 冲突）")
                    occ[rr][cc] = (r, c)
            origins.append((r, c, cell, r1, c1))
            c += cs
    for r in range(nrow):
        for c in range(ncol):
            if occ[r][c] is None:
                raise ValueError(
                    f"表格 {name} 的 ({r},{c}) 没有任何单元格覆盖 —— "
                    f"该行列数不足，请补齐 cells[{r}] 或检查 rowspan")

    cw = _resolve_dims(spec.get("col_widths"), wpx, ncol, "col_widths", name)
    rh = _resolve_dims(spec.get("row_heights"), hpx, nrow, "row_heights", name)

    gf = shapes.add_table(nrow, ncol, Emu(int(round(x1 * S))),
                          Emu(int(round(TOP_EMU + y1 * S))),
                          Emu(max(1, int(round(wpx * S)))), Emu(max(1, int(round(hpx * S)))))
    gf.name = name
    tbl = gf.table

    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for k in ("firstRow", "lastRow", "firstCol", "lastCol", "bandRow", "bandCol"):
            tblPr.set(k, "0")                      # 关掉主题表头/镶边
        for e in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(e)                        # 去掉主题样式 id

    for i, w in enumerate(cw):
        tbl.columns[i].width = Emu(int(round(w * S)))
    for i, h in enumerate(rh):
        tbl.rows[i].height = Emu(int(round(h * S)))

    # ---- 先合并，再写内容（merge 会清掉被并单元格的文字） ----
    # origins 里的合并区互不重叠（上面已校验），所以合并顺序无所谓。
    for (r, c, cell, r1, c1) in origins:
        if (r1, c1) != (r, c):
            tbl.cell(r, c).merge(tbl.cell(r1, c1))

    # ---- 内容 + 样式 ----
    hdr = int(spec.get("header_rows", 0) or 0)
    zebra = st.get("zebra_fill")
    for (r, c, cell, r1, c1) in origins:
        eff = dict(cell)
        if r < hdr:
            for k, v in (("bold", st.get("header_bold", True)),
                         ("color", st.get("header_color")),
                         ("fill", st.get("header_fill"))):
                if k not in eff and v is not None:
                    eff[k] = v
        if "fill" not in eff:
            if zebra is not None:
                z = (r - hdr) if hdr else r
                if z >= 0 and z % 2 == 1:
                    eff["fill"] = zebra
            if "fill" not in eff and st.get("cell_fill") is not None:
                eff["fill"] = st["cell_fill"]

        style_cell(tbl.cell(r, c), eff, st, write_text=True)
        if (r1, c1) != (r, c):                     # 被并单元格：只带样式，不写文字
            for rr in range(r, r1 + 1):
                for cc in range(c, c1 + 1):
                    if (rr, cc) != (r, c):
                        style_cell(tbl.cell(rr, cc), eff, st, write_text=False)
    return gf


# ---------------------------------------------------------------- z-order


def _has_text(el):
    tx = el.find(qn("p:txBody"))
    if tx is None:
        return False
    for r in tx.iter(qn("a:r")):
        t = r.find(qn("a:t"))
        if t is not None and (t.text or "").strip():
            return True
    return False


def place_z(spTree, elems, mode):
    """把 elems 按给定顺序移到目标层级。under-text = 第一个「有文字的 sp」之前。"""
    for el in elems:
        spTree.remove(el)
    idx = len(spTree)
    if mode == "under-text":
        for i, ch in enumerate(spTree):
            if ch.tag == qn("p:sp") and _has_text(ch):
                idx = i
                break
    for j, el in enumerate(elems):
        spTree.insert(idx + j, el)


# ---------------------------------------------------------------- probe


def probe(ref, box, dark=None, frac=0.5):
    """从原图探测表格网格线，给出 col_widths / row_heights 与底色建议。只读。

    dark=None 时用自适应阈值 `clip(中位灰度 - 20, 150, 245)`：
    表格底色常常是浅灰（实测网格线灰度 211、表头填充 241、白底 255），
    固定阈值 200 会整条漏掉；而「中位 - 20」能自动落在两者之间。
    """
    im = np.array(Image.open(ref).convert("RGB"))
    H, W = im.shape[:2]
    x1, y1, x2, y2 = [int(round(v)) for v in box]
    sub = im[y1:y2, x1:x2]
    g = sub.mean(axis=2)
    if dark is None:
        dark = float(max(150.0, min(245.0, float(np.median(g)) - 20.0)))
        print(f"自适应网格线阈值 dark={dark:.1f}（中位灰度 {np.median(g):.0f} - 20）")
    darkm = g < dark

    def peaks(score, lo, hi):
        # score 是「该列/行上暗像素的占比」（0~1），所以直接和 frac 比。
        m = score > frac
        out, run = [], []
        for i, v in enumerate(m):
            if v:
                run.append(i)
            elif run:
                out.append(int(round(sum(run) / len(run))) + lo)
                run = []
        if run:
            out.append(int(round(sum(run) / len(run))) + lo)
        return [v for v in out if lo + 2 < v < hi - 2]

    ex = sorted([x1] + peaks(darkm.mean(axis=0), x1, x2) + [x2])
    ey = sorted([y1] + peaks(darkm.mean(axis=1), y1, y2) + [y2])

    print(f"参考图 {ref}  {W}x{H}")
    print(f"探测框 [{x1},{y1},{x2},{y2}]  内部竖线 {len(ex)-2} 条  横线 {len(ey)-2} 条")
    print(f"  竖线 x = {ex}")
    print(f"  横线 y = {ey}")
    print(f"  col_widths = {[ex[i+1]-ex[i] for i in range(len(ex)-1)]}   // {len(ex)-1} 列")
    print(f"  row_heights = {[ey[i+1]-ey[i] for i in range(len(ey)-1)]}   // {len(ey)-1} 行")

    def mode_color(region):
        q = (region // 8).reshape(-1, 3)
        top = Counter(map(tuple, q)).most_common(1)[0][0]
        return [int(v) * 8 + 4 for v in top]

    ring = [im[y1:min(y2, y1 + 3), x1:x2].reshape(-1, 3),
            im[max(y1, y2 - 3):y2, x1:x2].reshape(-1, 3),
            im[max(0, y1 - 4):min(H, y2 + 4), max(0, x1 - 3):x1].reshape(-1, 3),
            im[max(0, y1 - 4):min(H, y2 + 4), x2:min(W, x2 + 3)].reshape(-1, 3)]
    print(f"  框内众数色 = {mode_color(sub)}      // cell_fill / cover.fill 的候选")
    print(f"  框外环带众数色 = {mode_color(np.vstack(ring))}  // cover.fill 的候选（表格外底色）")
    if len(ex) == 2 and len(ey) == 2:
        print("  ! 未探到网格线：可能是无框线表格，或阈值/框位不对"
              "（改用 --dark 手指定阈值，或收紧 box；--frac 调小可放宽）")


# ---------------------------------------------------------------- main


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd")
    pb = sub.add_parser("probe", help="从原图探测表格网格线（只读）")
    pb.add_argument("--ref", required=True)
    pb.add_argument("--box", required=True, help="x1,y1,x2,y2（画布像素坐标）")
    pb.add_argument("--dark", type=float, default=None,
                    help="网格线灰度阈值；缺省按框内中位灰度 - 20 自适应（clamp 到 150~245）")
    pb.add_argument("--frac", type=float, default=0.5, help="沿线命中比例阈值，缺省 0.5")

    ap.add_argument("--layout", default=None, help="layout.json（取 canvas / slide）")
    ap.add_argument("--canvas", default=None, help="WxH，配合 --slide 时可替代 --layout")
    ap.add_argument("--slide", default=None, help="WxH(pt)，缺省 960x540")
    ap.add_argument("--tables", default=None)
    ap.add_argument("--pptx", default=None, help="输入 PPTX（来自 layout_to_pptx.py）")
    ap.add_argument("--out", default=None)
    ap.add_argument("--z", default="under-text", choices=["under-text", "top"])
    ap.add_argument("--no-cover", action="store_true", help="忽略 spec 里的 cover，不画底板")
    ap.add_argument("--verify", action="store_true", help="写出后读回校验并打印结构")
    args = ap.parse_args()

    if args.cmd == "probe":
        probe(args.ref, [float(v) for v in args.box.split(",")], args.dark, args.frac)
        return

    for k in ("tables", "pptx", "out"):
        if not getattr(args, k):
            ap.error(f"build 模式必须给 --{k}")
    if not args.layout and not (args.canvas and args.slide):
        ap.error("必须给 --layout，或同时给 --canvas 与 --slide")

    if args.layout:
        base = os.path.dirname(os.path.abspath(args.layout))
        L = json.load(open(args.layout, encoding="utf-8"))
        CW, CH = L["canvas"]
        w_pt, h_pt = L["slide"]["w_pt"], L["slide"]["h_pt"]
    else:
        base = os.path.dirname(os.path.abspath(args.tables))
        CW, CH = [int(v) for v in args.canvas.lower().split("x")]
        w_pt, h_pt = [float(v) for v in args.slide.lower().split("x")]

    S = w_pt * EMU_PER_PT / CW
    TOP_EMU = (h_pt * EMU_PER_PT - CH * S) / 2.0
    print(f"canvas {CW}x{CH}  slide {w_pt:g}x{h_pt:g}pt  1px={S:.2f}emu")

    T = json.load(open(args.tables, encoding="utf-8"))
    tables = T["tables"] if isinstance(T, dict) else T
    st_global = (T.get("style") if isinstance(T, dict) else None) or {}

    prs = Presentation(args.pptx)
    slide = prs.slides[0]
    shapes = slide.shapes
    spTree = shapes._spTree
    n0 = len(spTree)

    for spec in tables:
        st = dict(st_global)
        st.update(spec.get("style") or {})
        name = spec.get("name") or "table"
        to_move = []

        cov = spec.get("cover")
        if cov and not args.no_cover:
            pad = float(cov.get("pad", 3))
            x1, y1, x2, y2 = spec["box"]
            sh = shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(int(round((x1 - pad) * S))), Emu(int(round(TOP_EMU + (y1 - pad) * S))),
                Emu(max(1, int(round((x2 - x1 + 2 * pad) * S)))),
                Emu(max(1, int(round((y2 - y1 + 2 * pad) * S)))))
            sh.name = name + "-cover"
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(*col(cov.get("fill"), st.get("cell_fill", [255, 255, 255])))
            sh.line.fill.background()
            sh.shadow.inherit = False
            to_move.append(sh._element)

        gf = build_table(shapes, spec, S, TOP_EMU, st)
        to_move.append(gf._element)

        mode = spec.get("z") or args.z
        place_z(spTree, to_move, mode)

        nrow, ncol = len(spec["cells"]), max(len(r) for r in spec["cells"])
        nm = sum(1 for r in spec["cells"] for c in r
                 if (int((c or {}).get("rowspan", 1) or 1) > 1
                     or int((c or {}).get("colspan", 1) or 1) > 1))
        print(f"  + {name}: {nrow}行 x {ncol}列  box={spec['box']}  "
              f"跨格 {nm}  表头 {spec.get('header_rows', 0)} 行  z={mode}")

    prs.save(args.out)
    print(f"已写出 {args.out}  (spTree {n0} -> {len(spTree)})")

    if args.verify:
        chk = Presentation(args.out)
        spTree = chk.slides[0].shapes._spTree
        for s in chk.slides[0].shapes:
            if not s.has_table:
                continue
            t = s.table
            merges = []
            for r in range(len(t.rows)):
                for c in range(len(t.columns)):
                    cell = t.cell(r, c)
                    if cell.is_merge_origin and (cell.span_height > 1 or cell.span_width > 1):
                        merges.append(f"({r},{c}){cell.span_height}x{cell.span_width}")
            print(f"  校验 {s.name}: {len(t.rows)}x{len(t.columns)}  "
                  f"行高pt={[round(x.height/EMU_PER_PT,1) for x in t.rows]}  "
                  f"列宽pt={[round(x.width/EMU_PER_PT,1) for x in t.columns]}")
            print(f"        合并区域 {merges or '无'}；"
                  f"有文字单元 {sum(1 for r in range(len(t.rows)) for c in range(len(t.columns)) if t.cell(r,c).text.strip())} 个")
        order = []
        for ch in spTree:
            if ch.tag == qn("p:pic"):
                order.append("pic")
            elif ch.tag == qn("p:graphicFrame"):
                order.append("table")
            elif ch.tag == qn("p:sp"):
                order.append("text" if _has_text(ch) else "shape")
        print(f"  层级（底→顶）: {' -> '.join(order)}")


if __name__ == "__main__":
    main()
